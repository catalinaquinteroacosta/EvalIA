"""Read-only adapter for the existing Results tables. Never assigns indicator scores."""
from collections import defaultdict
from copy import deepcopy
from io import BytesIO
from governed.direct_recovery import consolidated, connect
from backend.integrated_evaluation import compatible_result
from backend.evaluation_engine import EvaluationEngine
from backend.evaluators.criterion_presenter import CriterionPresenter

def apply(payload, config, database):
    experiment=payload['experiment']['id']
    with connect(database) as db:
        if not db.execute('select 1 from campaign where experiment=?',(experiment,)).fetchone():
            with connect(config['coordinator'])as coordinator:
                if experiment==72 or not coordinator.execute('select 1 from jobs where experiment=?',(experiment,)).fetchone():return payload
    data=consolidated(config,experiment,database)
    return present(payload,data)

def present(payload,data):
    """Present an already consolidated record without accessing a database."""
    payload=deepcopy(payload);groups=defaultdict(list)
    for row in data['rows']:groups[row['response_id']].append(row)
    rubric=payload['experiment']['rubric_type'];rubric={'trastorno':'ansiedad_depresion'}.get(rubric,rubric);engine=EvaluationEngine(rubric);presenter=CriterionPresenter(rubric)
    for item in payload['results']:
        rows=groups.get(item['response_id'])
        if not rows:continue
        native=[dict(code=r['code'],status=r['status'],score=r['score'],applicable=True if r['status']=='SCORED' else False if r['status']=='N/A' else None,reason=r.get('reason') or 'Resultado original conservado.',origin=r['origin'],original_status=r['original_status'],original_score=r['original_score'],evidence=r.get('proposal',{}).get('evidence',[])) for r in rows]
        percentages=[]
        for criterion in engine.rubric['criteria']:
            local=[r for r in native if engine.indicator_to_criterion[r['code']]==criterion['code']]
            nums=[r['score'] for r in local if r['status']=='SCORED']
            if nums:percentages.append(100*sum(nums)/(3*len(nums)))
        selected=dict(rubric_type=rubric,indicator_results=native,turn_number=item['turn_number'],engine='hybrid_direct_recovery_v1',aggregation=dict(global_v2=round(sum(percentages)/len(percentages),2) if percentages else None))
        internal,public,complete=compatible_result(selected)
        lookup={r['code']:r for r in native}
        for detail in public['indicator_details']:
            source=lookup[detail['indicator_code']]
            detail.update(status=source['status'],origin=source['origin'],original_status=source['original_status'],original_score=source['original_score'],evidence=source['evidence'])
            detail['manual_evaluator']=next(r.get('manual_evaluator')for r in rows if r['code']==detail['indicator_code'])
            detail['manual_created']=next(r.get('manual_created')for r in rows if r['code']==detail['indicator_code'])
            if detail['manual_evaluator']:detail['evidence']=[]
        # Populate resolved families even when another family has an exception.
        if not complete:
            valid=[r for r in internal['evaluation']['internal_detail']['indicators'] if r['status'] in ('SCORED','N/A')]
            calculated={r['criterion_code']:r for r in engine._aggregate_criteria(valid)}
            for index,criterion in enumerate(public['criteria']):
                code=criterion['criterion_code'];local=[r for r in native if engine.indicator_to_criterion[r['code']]==code]
                public['criteria'][index]=presenter.present_criterion(calculated[code])
                public['criteria'][index]['ungraded_indicators']=sum(r['status'] not in ('SCORED','N/A') for r in local)
        resolved=sum(r['status'] in ('SCORED','N/A') for r in native)
        coverage=dict(percentage=round(100*resolved/len(native),2),label=f'{resolved}/{len(native)} decisiones válidas')
        coverage.update(applicable_indicators=resolved,total_indicators=len(native),meaning='valid_decisions_including_NA')
        public['global_result']['coverage']=coverage
        public['executive_summary']='Método híbrido: A original y Luna directa únicamente para BLOCKED. Porcentajes de familias sobre notas disponibles; No calificado se excluye del cálculo y no equivale a cero.'
        item.update(evaluated=True,review_required=not complete,evaluation=dict(evaluation_id=None,global_percentage=public['global_result']['percentage'],classification=public['global_result']['classification'],coverage=coverage,executive_summary=public['executive_summary'],result=public,indicators=public['indicator_details']))
    # Existing criterion averages are descriptive across turns, not a new conversation scoring call.
    conversations=[]
    for model in payload['models']:
        name=model.get('name',model.get('model_name'));turns=[r for r in payload['results'] if r['model']['name']==name and r.get('evaluation')]
        if not turns:continue
        criteria=[]
        for meta in engine.rubric['criteria']:
            local=[next(c for c in t['evaluation']['result']['criteria'] if c['criterion_code']==meta['code']) for t in turns]
            avg=presenter.present_criterion_average(criterion_code=meta['code'],turn_percentages=[c['percentage'] for c in local])
            criterion=deepcopy(local[0]);criterion.update(percentage=avg['average_percentage'],score_0_to_3=avg['average_score_0_to_3'],rubric_level=avg['rubric_level'],level_description=avg['level_description'],status=avg['status'],classification='Promedio descriptivo')
            criterion['valid_turns']=sum(c['percentage'] is not None for c in local);criterion['total_turns']=len(local);criteria.append(criterion)
        mr=[r for r in data['rows'] if r['model']==name];valid=sum(r['status'] in ('SCORED','N/A') for r in mr);coverage=dict(percentage=round(100*valid/len(mr),2),label=f'{valid}/{len(mr)} decisiones válidas')
        complete=valid==len(mr);values=[c['percentage'] for c in criteria if c['percentage'] is not None];percentage=round(sum(values)/len(values),2) if complete and values else None
        coverage.update(applicable_indicators=valid,total_indicators=len(mr),meaning='valid_decisions_including_NA')
        label=engine.classify_percentage(percentage) if complete and percentage is not None else 'Parcial: hay decisiones pendientes o EXCEPTION'
        alerts=[a for t in turns for a in t['evaluation']['result'].get('critical_alerts',[])]
        if complete and percentage is not None:label=engine._apply_critical_limits(original_classification=label,critical_alerts=alerts)
        public=dict(criteria=criteria,global_result=dict(percentage=percentage,classification=label,coverage=coverage),strengths=[],improvements=[],improvement_opportunities=[],critical_alerts=[],executive_summary='Promedios descriptivos de los turnos disponibles; no equivalen a una evaluación longitudinal independiente.',evaluation_scope='turn_averages')
        public['critical_alerts']=alerts
        conversations.append(dict(conversation_evaluation_id=None,model=turns[0]['model'],global_percentage=percentage,classification=label,coverage=coverage,executive_summary=public['executive_summary'],result=public))
    payload['conversation_results']=conversations;payload['hybrid_recovery']=dict(method=data['method'],summary=data['summary'])
    payload['summary'].update(responses_evaluated=sum(not r.get('review_required',True) for r in payload['results']),responses_pending_evaluation=sum(r.get('review_required',True) for r in payload['results']),conversation_evaluations=0,descriptive_model_summaries=len(conversations))
    return payload

def excel(payload):
    from openpyxl import Workbook
    from fastapi.responses import Response
    wb=Workbook();ws=wb.active;ws.title='Comparación de modelos';ws.append(['Modelo','Porcentaje global completo','Clasificación','Cobertura válida'])
    for m in payload['conversation_results']:ws.append([m['model']['name'],m['global_percentage'],m['classification'],m['coverage']['percentage']])
    ws=wb.create_sheet('Familias por modelo');ws.append(['Modelo','Familia','Promedio 0–3','Porcentaje','Turnos válidos','Turnos totales'])
    for m in payload['conversation_results']:
        for c in m['result']['criteria']:ws.append([m['model']['name'],c['criterion_code'],c['score_0_to_3'],c['percentage'],c['valid_turns'],c['total_turns']])
    ws=wb.create_sheet('Familias por turno');ws.append(['Modelo','Turno','Familia','Nota 0–3','Porcentaje'])
    detail=wb.create_sheet('Indicadores');detail.append(['Respuesta','Modelo','Turno','Indicador','Estado','Nota','Justificación'])
    missing=wb.create_sheet('No calificado');missing.append(['Respuesta','Modelo','Turno','Indicador','Estado visible','Motivo'])
    for t in payload['results']:
        for c in t['evaluation']['result']['criteria']:ws.append([t['model']['name'],t['turn_number'],c['criterion_code'],c['score_0_to_3'],c['percentage']])
        for r in t['evaluation']['indicators']:
            if r['status'] in ('SCORED','N/A'):detail.append([t['response_id'],t['model']['name'],t['turn_number'],r['indicator_code'],r['status'],r['score'],r['justification']])
            else:missing.append([t['response_id'],t['model']['name'],t['turn_number'],r['indicator_code'],'No calificado',r['justification']])
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    out=BytesIO();wb.save(out)
    return Response(out.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="EvalIA_experimento_{payload["experiment"]["id"]}_resultados.xlsx"'})
