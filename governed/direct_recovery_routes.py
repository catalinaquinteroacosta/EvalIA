"""Experiment-owned hybrid results, exports, controlled recovery launch."""
import hashlib,json,subprocess,sys,time
from decimal import Decimal
from pathlib import Path
from direct_recovery import initialize,prepare,consolidated,csv_bytes,connect,wire,digest

def mount(main,config):
 from fastapi import HTTPException,Request
 from fastapi.responses import HTMLResponse,Response
 from governed_parallel import GovernedLab
 from queue_lab import units
 root=Path(config['project_root'])/'.evalia-private/direct_recovery';db=root/'recovery.db';initialize(db)
 main.app.state.direct_recovery_database=str(db)
 main.app.state.direct_recovery_config=config
 from governed.manual_grading import mount as mount_manual
 mount_manual(main,config,db)
 @main.app.get('/api/experiments/{experiment}/hybrid-results')
 def results(experiment:int):
  try:
   data=consolidated(config,experiment,db);lab=GovernedLab(config['coordinator'])
   with lab.connect()as c:
    j=c.execute('select * from jobs where experiment=?',(experiment,)).fetchone()
    used=c.execute('select coalesce(sum(actual),0)from requests where job=?',(j['id'],)).fetchone()[0];held=c.execute("select coalesce(sum(reserve),0)from requests where job=? and status in ('reserved','sending','uncertain')",(j['id'],)).fetchone()[0];old=c.execute('select coalesce(sum(spent),0),coalesce(sum(held),0)from legacy_commitments where job=?',(j['id'],)).fetchone()
    globalcap=c.execute('select cap from settings').fetchone()[0]
    globalcommitted=c.execute("select coalesce(sum(actual+case when status in ('reserved','sending','uncertain')then reserve else 0 end),0)from requests").fetchone()[0]+c.execute('select coalesce(sum(spent+held),0)from legacy_commitments').fetchone()[0]
    data['budget']=dict(spent_usd=str(Decimal(used+old[0])/10**9),reserved_usd=str(Decimal(held+old[1])/10**9),remaining_usd=str(Decimal(j['cap']-used-old[0]-held-old[1])/10**9),cap_usd=str(Decimal(j['cap'])/10**9),global_cap_usd=str(Decimal(globalcap)/10**9),worker_status=j['status'],warning='Saldo bajo'if j['cap']-used-old[0]-held-old[1]<100000000 else '')
    from monetary_policy import monitoring
    monitor_only=monitoring(c)
   data['budget']['global_remaining_usd']=str(Decimal(globalcap-globalcommitted)/10**9)
   data['budget']['mode']='monitor_only'if monitor_only else'enforced'
   if monitor_only:data['budget'].update(remaining_usd=None,global_remaining_usd=None,cap_usd=None,global_cap_usd=None,warning='Sin corte monetario local; consumo monitorizado')
   data['execution_state']='complete'if data['complete']else 'complete_with_exceptions'if not data['summary']['states'].get('BLOCKED',0)else 'pending'
   authorization=root/str(experiment)/'authorization.json'
   if authorization.exists():
    auth=json.loads(authorization.read_text(encoding='utf8'))
    with connect(db)as c:finished=c.execute("select count(*)from progress where campaign=? and state='completed' and updated>=?",(digest(json.loads((authorization.parent/'manifest.json').read_text(encoding='utf8'))),auth['at'])).fetchone()[0]
    elapsed=max(0,time.time()-auth['at']);remaining=max(0,auth['max_calls']-finished)
    data['timing']=dict(new_calls_total=auth['max_calls'],new_calls_finished=finished,elapsed_seconds=elapsed,estimated_remaining_seconds=elapsed*remaining/finished if finished else None,estimate_is_approximate=True)
   return data
  except ValueError as e:raise HTTPException(409,str(e))
 @main.app.get('/api/experiments/{experiment}/hybrid-results.csv')
 def export(experiment:int):return Response(csv_bytes(results(experiment)),media_type='text/csv; charset=utf-8',headers={'Content-Disposition':f'attachment; filename="EvalIA_{experiment}_hibrido.csv"'})
 @main.app.get('/api/experiments/{experiment}/hybrid-results.json')
 def export_json(experiment:int):return Response(wire(results(experiment)),media_type='application/json',headers={'Content-Disposition':f'attachment; filename="EvalIA_{experiment}_hibrido.json"'})
 @main.app.get('/api/experiments/{experiment}/hybrid-results.xlsx')
 def export_xlsx(experiment:int):
  # Application export using the project's existing Excel engine, not a scorer.
  from io import BytesIO
  from openpyxl import Workbook
  from openpyxl.styles import Font,PatternFill,Alignment
  data=results(experiment);book=Workbook();meta=book.active;meta.title='Resumen'
  for row in [('Experimento',experiment),('Método',data['method']),('Estado','Completo'if data['complete']else 'Con pendientes'),('Total indicadores',data['summary']['total']),('Notas numéricas',data['summary']['numeric']),('Cobertura (%)',data['summary']['coverage']),('Estados',wire(data['summary']['states'])),('Gasto confirmado USD',float(data['budget']['spent_usd'])),('Reservado USD',float(data['budget']['reserved_usd'])),('Disponible USD',float(data['budget']['remaining_usd'])),('Criterio de porcentaje',data['note'])]:meta.append(row)
  for title,key,dimension in [('Modelo y turno','by_model_turn','turn'),('Modelo y familia','by_model_family','family')]:
   sheet=book.create_sheet(title);sheet.append(['Modelo','Turno'if dimension=='turn'else 'Familia','Total','Numéricos','Suma notas','Promedio 0–3','Porcentaje disponible','Cobertura (%)','Estados'])
   for r in data[key]:sheet.append([r['model'],r[dimension],r['total'],r['numeric'],r['score_sum'],r['mean'],r['percentage'],r['coverage'],wire(r['states'])])
  sheet=book.create_sheet('Indicadores');sheet.append(['Respuesta','Modelo','Turno','Indicador','Estado original','Nota original','Estado consolidado','Nota consolidada','Origen','Justificación recuperación'])
  for r in data['rows']:sheet.append([r['response_id'],r['model'],r['turn'],r['code'],r['original_status'],r['original_score'],r['status'],r['score'],r['origin'],r.get('reason','')])
  for sheet in book:
   sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
   for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='203D61')
   for col in sheet.columns:
    letter=col[0].column_letter;sheet.column_dimensions[letter].width=min(55,max(15,max(len(str(c.value or ''))for c in list(col)[:50])+2))
   for row in sheet:
    for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
  output=BytesIO();book.save(output)
  return Response(output.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="EvalIA_{experiment}_hibrido.xlsx"'})
 @main.app.post('/api/experiments/{experiment}/direct-recovery/prepare')
 def plan(experiment:int):
  try:
   cid,m=prepare(config,experiment,db,root/str(experiment))
   with connect(db)as c:done={r['id']for r in c.execute('select id from decisions where campaign=?',(cid,))}
   pending=[r for r in m['requests']if r['id']not in done]
   return dict(campaign=cid,experiment=experiment,max_calls=len(pending),estimate_usd=str(sum(Decimal(r['reserved_usd'])for r in pending)),reused=len(done),paid_execution_started=False)
  except ValueError as e:raise HTTPException(409,str(e))
 @main.app.post('/api/experiments/{experiment}/direct-recovery/confirm')
 async def confirm(experiment:int,request:Request):
  from urllib.parse import urlsplit
  if request.headers.get('origin')and urlsplit(request.headers['origin']).hostname not in ('127.0.0.1','localhost','::1'):raise HTTPException(403,'Solo confirmación local')
  d=await request.json()
  if d.get('confirmation')!='AUTORIZO':raise HTTPException(403,'Requiere autorización del lote y presupuesto')
  folder=root/str(experiment);file=folder/'manifest.json'
  if not file.exists():raise HTTPException(409,'Preparar primero el lote')
  m=json.loads(file.read_text(encoding='utf8'));cid=digest(m)
  if d.get('campaign')!=cid:raise HTTPException(409,'El lote cambió')
  if (folder/'authorization.json').exists():raise HTTPException(409,'Lote autorizado; no duplicar llamadas')
  lab=GovernedLab(config['coordinator'])
  with connect(db)as c:done={r['id']for r in c.execute('select id from decisions where campaign=?',(cid,))}
  rows=[r for r in m['requests']if r['id']not in done];estimate=sum(Decimal(r['reserved_usd'])for r in rows)
  cap=units(d.get('individual_cap_usd','0'));globalcap=units(d.get('global_cap_usd','0'))
  auth=dict(explicit_user_approval=True,manifest_sha256=hashlib.sha256(file.read_bytes()).hexdigest(),allowed_ids=[r['id']for r in rows],max_calls=len(rows),individual_cap_usd=str(Decimal(cap)/10**9),global_cap_usd=str(Decimal(globalcap)/10**9),no_uncertain_retries=True,at=time.time())
  with lab.tx()as c:
   j=c.execute('select * from jobs where experiment=?',(experiment,)).fetchone()
   if not j or j['owner']is not None:raise HTTPException(409,'Experimento ocupado')
   auth['job']=j['id'];currentglobal=c.execute('select cap from settings').fetchone()[0]
   def committed(where='',args=()):return c.execute("select coalesce(sum(actual+case when status in ('reserved','sending','uncertain')then reserve else 0 end),0)from requests "+where,args).fetchone()[0]+c.execute('select coalesce(sum(spent+held),0)from legacy_commitments '+where,args).fetchone()[0]
   if cap<j['cap']or globalcap<currentglobal:raise HTTPException(422,'No reducir límites existentes')
   from monetary_policy import monitoring
   if not monitoring(c) and (committed('where job=?',(j['id'],))+units(estimate)>cap or committed()+units(estimate)>globalcap):raise HTTPException(422,'Presupuesto insuficiente para el lote completo')
   for scope,scopeid,old,new in [('individual',j['id'],j['cap'],cap),('global','global',currentglobal,globalcap)]:
    if old!=new:c.execute('insert into budget_amendments values(?,?,?,?,?,?,?,?)',(cid+'-'+scope,scope,scopeid,old,new,wire(auth),digest(auth),time.time()))
   c.execute('update jobs set cap=? where id=?',(cap,j['id']));c.execute('update settings set cap=? where id=1',(globalcap,));position=c.execute('select max(position)from requests where job=?',(j['id'],)).fetchone()[0]+1
   for i,r in enumerate(rows):c.execute('insert into requests(key,job,position,body,reserve,status)values(?,?,?,?,?,?)',(digest(r['body']),j['id'],position+i,wire(r['body']),units(r['reserved_usd']),'ready'))
  (folder/'authorization.json').write_text(wire(auth),encoding='utf8')
  cfgpath=folder/'config.json';cfgpath.write_text(wire(config),encoding='utf8')
  log=(folder/'worker.log').open('ab')
  try:
   proc=subprocess.Popen([config['python'],'-X','utf8','-B',str(Path(__file__).with_name('direct_recovery_worker.py')),'--config',str(cfgpath),'--campaign',cid,'--database',str(db),'--folder',str(folder)],cwd=config['project_root'],stdout=log,stderr=log,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
  finally:log.close()
  return dict(started=True,campaign=cid,pid=proc.pid,max_calls=len(rows),estimate_usd=str(estimate))
 @main.app.post('/api/experiments/{experiment}/direct-recovery/resume')
 async def resume(experiment:int,request:Request):
  from urllib.parse import urlsplit
  from supervision import Fence
  if request.headers.get('origin')and urlsplit(request.headers['origin']).hostname not in ('127.0.0.1','localhost','::1'):raise HTTPException(403,'Solo operación local')
  folder=root/str(experiment);file=folder/'authorization.json'
  if not file.exists():raise HTTPException(409,'No hay lote previamente autorizado')
  auth=json.loads(file.read_text(encoding='utf8'));fence=Fence(Path(config['fences'])/(auth['job']+'.worker.lock'))
  if not fence.acquire():raise HTTPException(409,'El trabajador sigue activo')
  lab=GovernedLab(config['coordinator'])
  try:
   with lab.tx()as c:
    # A crash after reservation is conservatively uncertain, never a free retry.
    c.execute("update requests set status='uncertain'where job=? and status in ('sending','reserved')",(auth['job'],))
    c.execute("update jobs set owner=NULL,status='stopped_requires_review'where id=?",(auth['job'],))
  finally:fence.release()
  m=json.loads((folder/'manifest.json').read_text(encoding='utf8'));log=(folder/'worker.log').open('ab')
  try:proc=subprocess.Popen([config['python'],'-X','utf8','-B',str(Path(__file__).with_name('direct_recovery_worker.py')),'--config',str(folder/'config.json'),'--campaign',digest(m),'--database',str(db),'--folder',str(folder)],cwd=config['project_root'],stdout=log,stderr=log,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
  finally:log.close()
  return dict(started=True,pid=proc.pid,policy='Only authorized ready cases; no uncertain retries')
 @main.app.get('/experimento-recuperacion',response_class=HTMLResponse)
 def panel():return Path(__file__).with_name('direct_recovery.html').read_text(encoding='utf8')
