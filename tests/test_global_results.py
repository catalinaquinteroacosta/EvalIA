from copy import deepcopy
import pytest
from backend.global_results import aggregate, workbook
from governed.double_resolution_v1.engine import catalog

def fixture_rows():
    experiments=[dict(id=1,code='ANS-01',category='Ansiedad')]
    rows=[]
    for turn in range(1,6):
        for code in catalog():
            if not code.startswith('CR-'):
                rows.append(dict(experiment=1,response_id=turn,model='ChatGPT',turn=turn,code=code,status='N/A',score=None))
    return experiments,rows

def test_equal_family_weight_zero_and_na():
    experiments,rows=fixture_rows()
    for row in rows:
        if row['code'] in ('AF1','AF2','CP1'):
            row.update(status='SCORED',score=0 if row['code'].startswith('AF') else 3)
    result=aggregate(experiments,rows)
    model=result['groups'][0]['models'][0]
    assert model['percentage']==50
    assert model['numeric']==15 and model['na']==390
    assert model['pending']==0 and model['coverage']==100
    assert result['groups'][0]['models'][1]['percentage'] is None
    assert result['groups'][0]['models'][1]['pending']==405

def test_missing_indicator_withholds_comparison():
    experiments,rows=fixture_rows()
    for row in rows:
        if row['code']=='AF1':row.update(status='SCORED',score=3)
    rows.pop(0)
    model=aggregate(experiments,rows)['groups'][0]['models'][0]
    assert model['percentage'] is None and model['pending']==1
    assert model['turns'][0]['percentage'] is None
    assert model['turns'][1]['percentage']==100

def test_rejects_duplicate_and_invalid_grade():
    experiments,rows=fixture_rows()
    with pytest.raises(ValueError,match='Duplicate'):aggregate(experiments,rows+[deepcopy(rows[0])])
    rows[0].update(status='SCORED',score=True)
    with pytest.raises(ValueError,match='Invalid'):aggregate(experiments,rows)

def test_combined_group_weights_scenarios_and_export_roundtrip():
    from io import BytesIO
    from openpyxl import load_workbook
    experiments,rows=fixture_rows()
    experiments += [dict(id=2,code='ANS-02',category='Ansiedad'),dict(id=3,code='DEP-01',category='Depresión')]
    base=deepcopy(rows)
    for exp in (2,3):
        for original in base:
            row=deepcopy(original);row['experiment']=exp;row['response_id']+=exp*10;rows.append(row)
    for row in rows:
        if row['code']=='AF1':row.update(status='SCORED',score=3 if row['experiment']==3 else 0)
    data=aggregate(experiments,rows)
    assert data['groups'][2]['models'][0]['percentage']==pytest.approx(100/3)
    assert data['groups'][2]['models'][0]['sd_scenarios']==pytest.approx(57.735026918962575)
    data.update(experiments=experiments,excluded=[],method=['Método'],generated_at='test',dataset_id='test',bank='test')
    book=load_workbook(BytesIO(workbook(data)))
    assert book['Trastornos']['B8'].value==pytest.approx(100/3)
    assert len(book.sheetnames)==12

def test_signature_ignores_bookkeeping_but_detects_grade_changes(tmp_path):
    import sqlite3
    from backend.global_results import signature
    paths=[tmp_path/f'{name}.db' for name in ('main','coordinator','manual')]
    for path in paths:
        with sqlite3.connect(path) as db:db.execute('CREATE TABLE bookkeeping(value)')
    with sqlite3.connect(paths[2]) as db:
        db.execute('CREATE TABLE manual_grades(experiment,response_id,code,score)')
        db.execute("INSERT INTO manual_grades VALUES(1,1,'AF1',0)")
    config=dict(database=str(paths[0]),coordinator=str(paths[1]))
    before=signature(config,paths[2])
    with sqlite3.connect(paths[1]) as db:db.execute('INSERT INTO bookkeeping VALUES(1)')
    assert signature(config,paths[2])==before
    with sqlite3.connect(paths[2]) as db:db.execute('UPDATE manual_grades SET score=1')
    assert signature(config,paths[2])!=before

def test_global_classification_uses_existing_rubric_without_mutating_grades():
    from backend.global_results import classified_report
    from backend.evaluation_engine import EvaluationEngine
    experiments,rows=fixture_rows()
    for row in rows:
        if row['code']=='AF1':row.update(status='SCORED',score=3)
    source=aggregate(experiments,rows)
    result=classified_report(source)
    assert result['groups'][0]['models'][0]['classification']==EvaluationEngine('ansiedad_depresion').classify_percentage(100)
    assert result['groups'][0]['models'][1]['classification']=='Sin resultado completo'
    assert 'classification' not in source['groups'][0]['models'][0]

def test_indicator_and_criticality_denominators():
    import json
    from pathlib import Path
    experiments,rows=fixture_rows()
    for row in rows:
        if row['code']=='AF1':row.update(status='SCORED',score=0 if row['turn']==1 else 3)
    group=aggregate(experiments,rows)['groups'][0]
    indicator=next(i for i in group['indicators'] if i['code']=='AF1')
    own=indicator['models'][0]
    assert own['percentage']==80 and own['numeric']==5 and own['zero']==1 and own['three']==4
    assert indicator['models'][1]['percentage'] is None and indicator['models'][1]['pending']==5
    rubric=json.loads((Path(__file__).resolve().parents[1]/'database/rubrics/ansiedad_depresion.json').read_text(encoding='utf8'))
    meta=next(i for c in rubric['criteria'] for i in c['indicators'] if i['code']=='AF1')
    assert indicator['criticality']==meta['criticality']
    criticality=next(c for c in group['criticalities'] if c['level']==meta['criticality'])['models'][0]
    assert criticality['percentage']==80 and criticality['numeric']==5
    assert criticality['na']>0 and criticality['zero']==1
    rows=[r for r in rows if not(r['code']=='AF1' and r['turn']==1)]
    group=aggregate(experiments,rows)['groups'][0]
    assert next(i for i in group['indicators'] if i['code']=='AF1')['models'][0]['percentage'] is None

def test_family_summary_includes_every_official_family_and_keeps_zero():
    from backend.global_results import classified_report
    experiments,rows=fixture_rows()
    for row in rows:
        if row['code']=='AF1':row.update(status='SCORED',score=0 if row['turn']==1 else 3)
    data=aggregate(experiments,rows);data['experiments']=experiments
    groups=classified_report(data)['groups']
    assert len(groups[0]['family_summaries'])==9
    assert len(groups[3]['family_summaries'])==10
    family=next(f for f in groups[0]['family_summaries'] if f['name']=='Afrontamiento')
    assert family['models'][0]['percentage']==80
    assert family['models'][0]['scenarios_numeric']==1
    assert family['models'][1]['percentage'] is None

def test_safety_matches_individual_engine_and_deduplicates_scenarios():
    from backend.global_results import classified_report
    from backend.evaluation_engine import EvaluationEngine
    experiments,rows=fixture_rows()
    engine=EvaluationEngine('ansiedad_depresion')
    codes=[code for code,meta in engine.indicators_by_code.items() if engine._is_critical_failure(meta,True,0)]
    assert len(codes)>=2
    for row in rows:
        if row['code'] in codes[:2] and row['turn'] in (1,2):row.update(status='SCORED',score=0 if row['turn']==1 else 1)
        elif row['code']==codes[0]:row.update(status='SCORED',score=2)
    data=aggregate(experiments,rows);data['experiments']=experiments
    assert len(data['safety_alerts'])==4
    summary=classified_report(data)['groups'][0]['safety'][0]
    assert summary['alerts']==4 and summary['scenarios_affected']==1
    assert summary['responses_affected']==2 and summary['responses_total']==5
    assert summary['complete'] is True
    absent=classified_report(data)['groups'][0]['safety'][1]
    assert absent['alerts']==0 and absent['complete'] is False
    native=[dict(indicator_code=r['code'],applicable=r['status']=='SCORED',score=r['score'],justification='Test') for r in rows if r['turn']==1]
    assert len(engine.validate_evaluation({'indicators':native})['critical_alerts'])==2
