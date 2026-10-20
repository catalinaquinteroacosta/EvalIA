import json
from io import BytesIO
from openpyxl import Workbook, load_workbook
from backend.export_audit import append_xlsx_audit, indicator_records


def payload():
    rows = [dict(code='CR-A1', status='SCORED', score=0, applicable=True),
            dict(code='CR-A2', status='N/A', score=None, applicable=False),
            dict(code='CR-A3', status='BLOCKED', score=None, applicable=None,
                 decisions={'p': {'extraction': {'citations': [{'quote': 'texto literal'}]},
                                  'verification': {'reason': 'contraevidencia'}}}, reason='x'*33000)]
    return dict(experiment=dict(id=999,scenario_title='Prueba',rubric_type='crisis'),
        results=[dict(response_id=1,turn_number=1,model=dict(id=1,name='Modelo'),evaluated=True,
            evaluation=dict(review_id=2,result=dict(engine_version='v-test',validated_audit=dict(
                reader_model='lector-test',indicator_results=rows))))])


def test_lossless_xlsx_null_and_review_without_longitudinal():
    book=Workbook(); book.active.title='Historica'; book.active['A1']='preservada'
    append_xlsx_audit(book,payload())
    output=BytesIO(); book.save(output); output.seek(0)
    loaded=load_workbook(output)
    assert loaded['Historica']['A1'].value=='preservada'
    rows=list(loaded['Indicadores por turno'].values)
    assert [row[4:7] for row in rows[1:]]==[('SCORED',0,True),('N/A',None,False),('BLOCKED',None,None)]
    pieces=[row[7] for row in list(loaded['Auditoria JSON'].values)[1:] if row[3]=='CR-A3']
    decoded=json.loads(''.join(pieces))
    assert decoded['score'] is None and len(decoded['reason'])==33000
    assert decoded['decisions']['p']['verification']['reason']=='contraevidencia'


def test_pdf_includes_review_without_conversation_and_literal_evidence():
    from backend.pdf_export import build_results_pdf
    from pypdf import PdfReader
    data=payload();data['results'][0]['evaluation']['result']['validated_audit']['indicator_results'][2]['reason']='Revisar'
    document=PdfReader(BytesIO(build_results_pdf(data)))
    text='\n'.join(page.extract_text() for page in document.pages)
    text=' '.join(text.split())
    assert 'CR-A3' in text and 'BLOCKED' in text and 'texto literal' in text
    assert 'contraevidencia' in text and 'review_id=2' in text and 'lector-test' in text
    assert 'Respuestas con revisión guardada' in text and 'sin resultado global longitudinal' in text


def test_pdf_removes_invisible_word_joiner_only_in_presentation():
    from backend.pdf_export import safe
    title='Une \u2060PAR-02\u2060 y PAR-03'
    assert safe(title)=='Une PAR-02 y PAR-03'
    assert '\u2060' in title


def test_anxiety_native_applicability_preserves_unresolved():
    data=payload()
    native=data['results'][0]['evaluation']['result']['validated_audit']['indicator_results']
    for row in native:
        row.pop('applicable')
    native[2]['applicability']='UNRESOLVED'
    assert [r['applicable'] for r in indicator_records(data)]==[True,False,None]
    native[2]['applicability']='APPLICABLE'
    assert indicator_records(data)[2]['applicable'] is True
