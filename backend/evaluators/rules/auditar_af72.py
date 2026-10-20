"""Audita AF1–AF8 del experimento 72 sin escribir en la base de datos."""

from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from backend.evaluators.evaluation_context import ConversationTurn, EvaluationContext
from backend.evaluators.rules.afrontamiento_rules import afrontamiento_rules


CODES = tuple(f"AF{i}" for i in range(1, 9))


def load_experiment(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Conversaciones"]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip().casefold() for value in next(rows)]

        def index_of(prefix: str) -> int | None:
            return next(
                (i for i, header in enumerate(headers) if header.startswith(prefix)),
                None,
            )

        columns = {
            "Turno": index_of("turno"),
            "Modelo": index_of("modelo"),
            "Mensaje del usuario": index_of("mensaje del usuario"),
            "Respuesta de la IA": index_of("respuesta"),
        }

        if any(index is None for index in columns.values()):
            raise ValueError(
                "La hoja Conversaciones necesita columnas de turno, modelo, "
                "mensaje del usuario y respuesta. "
                f"Encabezados encontrados: {headers}"
            )

        entries = [
            {label: row[index] for label, index in columns.items()}
            for row in rows
            if row[columns["Turno"]] is not None
        ]
    finally:
        workbook.close()

    if len(entries) != 25:
        raise ValueError(f"Se esperaban 25 respuestas, hay {len(entries)}.")

    if any(
        not str(row["Mensaje del usuario"] or "").strip()
        or not str(row["Respuesta de la IA"] or "").strip()
        for row in entries
    ):
        raise ValueError(
            "Falta un mensaje de usuario o una respuesta en Conversaciones."
        )

    keys = [(int(row["Turno"]), str(row["Modelo"])) for row in entries]
    if len(set(keys)) != 25 or set(t for t, _ in keys) != set(range(1, 6)):
        raise ValueError("Se esperaban cinco turnos sin duplicados por modelo.")

    return entries


def load_reference(path: Path) -> dict[tuple[int, str, str], int | None]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook["AF calificaciones"].iter_rows(values_only=True)
        headers = next(rows)
        result = {}
        for row in rows:
            data = dict(zip(headers, row))
            key = (int(data["Turno"]), str(data["Modelo"]), str(data["Indicador"]))
            grade = data["Puntaje 0–3"]
            result[key] = None if grade == "N/A" else int(grade)
    finally:
        workbook.close()

    if len(result) != 200:
        raise ValueError(
            f"La primera valoración AF debe contener 200 filas; hay {len(result)}."
        )
    return result


def evaluate(entries: list[dict], reference: dict | None) -> list[dict]:
    by_model = defaultdict(list)
    for row in entries:
        by_model[str(row["Modelo"])].append(row)

    results = []
    for model, model_rows in by_model.items():
        model_rows.sort(key=lambda row: int(row["Turno"]))
        history = []

        for row in model_rows:
            turn_number = int(row["Turno"])
            history.append(
                ConversationTurn(
                    turn_number=turn_number,
                    user_message=str(row["Mensaje del usuario"]),
                    ai_response=str(row["Respuesta de la IA"]),
                )
            )
            context = EvaluationContext(
                experiment_id=72,
                scenario_code="ANS-01",
                rubric_type="ansiedad_depresion",
                model_name=model,
                turns=list(history),
            )

            for code in CODES:
                verdict = afrontamiento_rules.evaluate(code, context, turn_number)
                key = (turn_number, model, code)
                score = verdict.score if verdict.applicable else None
                expected = reference[key] if reference is not None else None

                results.append(
                    {
                        "turno": turn_number,
                        "modelo": model,
                        "indicador": code,
                        "puntaje_motor": "N/A" if score is None else score,
                        "primera_valoracion": (
                            ""
                            if reference is None
                            else "N/A" if expected is None else expected
                        ),
                        "coincide_con_primera_valoracion": (
                            ""
                            if reference is None
                            else "SI" if score == expected else "NO"
                        ),
                        "version_regla": verdict.rule_version,
                        "justificacion_motor": verdict.justification,
                    }
                )

    return sorted(
        results,
        key=lambda row: (row["turno"], row["modelo"], row["indicador"]),
    )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "experimento",
        type=Path,
        help="Exportación Excel del experimento 72",
    )
    parser.add_argument(
        "--referencia",
        type=Path,
        help="Primera valoración AF (opcional; aún no es patrón definitivo)",
    )
    parser.add_argument(
        "--salida",
        type=Path,
        help="Archivo CSV de salida (opcional)",
    )
    args = parser.parse_args()

    reference = load_reference(args.referencia) if args.referencia else None
    rows = evaluate(load_experiment(args.experimento), reference)

    print("Experimento 72 — AF1 a AF8 (N/A es distinto de 0)")
    print("Turno Modelo       " + " ".join(f"{code:>3}" for code in CODES))

    for turn in range(1, 6):
        for model in ("ChatGPT", "Claude", "Copilot", "Gemini", "Perplexity"):
            values = [
                row
                for row in rows
                if row["turno"] == turn and row["modelo"] == model
            ]
            print(
                f"{turn:>5} {model:<12} "
                + " ".join(f"{str(item['puntaje_motor']):>3}" for item in values)
            )

    if reference is not None:
        print("\nCoincidencias con la PRIMERA valoración (25 por indicador):")
        for code in CODES:
            matches = sum(
                row["coincide_con_primera_valoracion"] == "SI"
                for row in rows
                if row["indicador"] == code
            )
            print(f"{code}: {matches}/25")

    if args.salida:
        with args.salida.open("w", newline="", encoding="utf-8-sig") as target:
            writer = csv.DictWriter(target, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
        print(f"\nCSV guardado: {args.salida}")


if __name__ == "__main__":
    main()