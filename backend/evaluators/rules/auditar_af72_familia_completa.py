"""
Auditoria integral de la familia AF sobre el Experimento 72.

Contrato:
    evaluate_af_family(...)
        -> result["indicators"]["AF1"..."AF8"]

Interpretacion:
    applicable is False -> N/A
    status == error     -> ERROR
    status == blocked   -> BLOCKED
    decision.score      -> 0, 1, 2 o 3
    decision ausente    -> UNRESOLVED

No modifica la base de datos.
"""

from pathlib import Path
import argparse
import csv
import json
from datetime import datetime, timezone
import time

from openpyxl import load_workbook

from backend.evaluators.af_family_evaluator import (
    evaluate_af_family,
)


TRACE_SCHEMA_VERSION = "1.0.0"


def trace_json_default(value):
    """Serializa estructuras conocidas sin ocultar tipos desconocidos."""
    from dataclasses import asdict, is_dataclass
    from enum import Enum

    if is_dataclass(value) and not isinstance(value, type):
        return asdict(value)

    if isinstance(value, Enum):
        return value.value

    raise TypeError(
        f"Tipo no serializable en trazabilidad: "
        f"{type(value).__module__}.{type(value).__qualname__}"
    )


INDICATORS = [
    "AF1",
    "AF2",
    "AF3",
    "AF4",
    "AF5",
    "AF6",
    "AF7",
    "AF8",
]


def indicator_block(result, indicator):

    if not isinstance(result, dict):
        return None

    indicators = result.get(
        "indicators"
    )

    if not isinstance(indicators, dict):
        return None

    block = indicators.get(
        indicator
    )

    if not isinstance(block, dict):
        return None

    return block


def score_value(result, indicator):

    block = indicator_block(
        result,
        indicator,
    )

    if block is None:
        return "ERROR"

    status = block.get(
        "status"
    )

    if status == "blocked":
        return "BLOCKED"

    if status == "error":
        return "ERROR"

    if block.get("error"):
        return "ERROR"

    applicable = block.get(
        "applicable"
    )

    if applicable is False:
        return "N/A"

    decision = block.get(
        "decision"
    )

    if decision is None:
        return "UNRESOLVED"

    if not isinstance(decision, dict):
        return "ERROR"

    value = decision.get(
        "score"
    )

    if value is None:
        return "UNRESOLVED"

    if value not in [
        0,
        1,
        2,
        3,
    ]:
        return "ERROR"

    return value


def technical_problem(value):

    return value in {
        "ERROR",
        "BLOCKED",
        "UNRESOLVED",
    }


def load_conversations(path):
    try:

        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
        )

        if "Conversaciones" not in workbook.sheetnames:
            raise ValueError(
                "No existe la hoja Conversaciones."
            )

        sheet = workbook["Conversaciones"]

        headers = [
            str(cell.value).strip()
            if cell.value is not None
            else ""
            for cell in sheet[1]
        ]

        required = [
            "Turno",
            "Modelo",
            "Mensaje del usuario",
            "Respuesta de la IA",
        ]

        missing = [
            name
            for name in required
            if name not in headers
        ]

        if missing:
            raise ValueError(
                f"Faltan columnas: {missing}"
            )

        idx = {
            name: headers.index(name)
            for name in required
        }

        rows = []

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True,
        ):

            user = row[
                idx["Mensaje del usuario"]
            ]

            ai = row[
                idx["Respuesta de la IA"]
            ]

            if user is None or ai is None:
                continue

            rows.append(
                {
                    "turno": row[
                        idx["Turno"]
                    ],
                    "modelo": str(
                        row[idx["Modelo"]]
                    ).strip(),
                    "user": str(user),
                    "ai": str(ai),
                }
            )

        return rows

    finally:
        if 'workbook' in locals():
            workbook.close()


def run_audit(
    input_path,
    output_path,
    model,
    timeout,
    limit=None,
):

    conversations = load_conversations(
        input_path
    )

    if limit is not None:
        conversations = conversations[
            :limit
        ]

    print("=" * 90)
    print(
        "EVALIA - GRAN PRUEBA FAMILIA AF / "
        "EXPERIMENTO 72"
    )
    print("=" * 90)

    print(
        "INPUT         :",
        input_path,
    )

    print(
        "OUTPUT        :",
        output_path,
    )

    print(
        "MODEL         :",
        model,
    )

    print(
        "CONVERSATIONS :",
        len(conversations),
    )

    print(
        "LIMIT         :",
        limit,
    )

    context = {}
    results = []
    traces = []

    start_total = time.perf_counter()

    for number, row in enumerate(
        conversations,
        start=1,
    ):

        modelo = row["modelo"]

        previous = context.get(
            modelo,
            {
                "user": "",
                "ai": "",
            },
        )

        print()
        print("=" * 90)

        print(
            f"[{number}/{len(conversations)}] "
            f"Turno={row['turno']} "
            f"Modelo={modelo}"
        )

        print("=" * 90)

        start = time.perf_counter()

        family = None

        try:

            family = evaluate_af_family(
                current_user=row["user"],
                current_ai=row["ai"],
                prior_user=previous["user"],
                prior_ai=previous["ai"],
                model=model,
                timeout=timeout,
            )

            values = {
                indicator: score_value(
                    family,
                    indicator,
                )
                for indicator in INDICATORS
            }

            exception = ""

        except Exception as exc:

            values = {
                indicator: "ERROR"
                for indicator in INDICATORS
            }

            exception = repr(exc)

        elapsed = (
            time.perf_counter()
            - start
        )

        technical = any(
            technical_problem(
                values[indicator]
            )
            for indicator in INDICATORS
        )

        print(
            " | ".join(
                f"{indicator}={values[indicator]}"
                for indicator in INDICATORS
            )
        )

        print(
            "TECHNICAL PROBLEM:",
            technical,
        )

        print(
            "TIME:",
            f"{elapsed:.2f}s",
        )

        if exception:
            print(
                "EXCEPTION:",
                exception,
            )

        result_row = {
            "Turno": row["turno"],
            "Modelo": modelo,
        }

        result_row.update(
            values
        )

        result_row[
            "TechnicalProblem"
        ] = technical

        result_row[
            "Exception"
        ] = exception

        result_row[
            "Seconds"
        ] = round(
            elapsed,
            3,
        )

        results.append(
            result_row
        )

        traces.append({
            "turno": row["turno"],
            "modelo": modelo,
            "evaluation_index": number,
            "scores": dict(values),
            "technical_problem": technical,
            "exception": exception,
            "elapsed_seconds": round(elapsed, 3),
            "family_result": family,
        })

        context[modelo] = {
            "user": row["user"],
            "ai": row["ai"],
        }

    trace_path = output_path.with_suffix(".trace.json")
    trace_document = {
        "trace_schema_version": TRACE_SCHEMA_VERSION,
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "input_filename": input_path.name,
        "model": model,
        "evaluations": traces,
    }

    total = (
        time.perf_counter()
        - start_total
    )

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with output_path.open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as fh:

        writer = csv.DictWriter(
            fh,
            fieldnames=[
                "Turno",
                "Modelo",
                *INDICATORS,
                "TechnicalProblem",
                "Exception",
                "Seconds",
            ],
        )

        writer.writeheader()
        writer.writerows(
            results
        )

    with trace_path.open(
        "w",
        encoding="utf-8",
    ) as fh:
        json.dump(
            trace_document,
            fh,
            ensure_ascii=False,
            indent=2,
            default=trace_json_default,
        )

    print("TRACE JSON    :", trace_path)

    problems = [
        row
        for row in results
        if row["TechnicalProblem"]
    ]

    print()
    print("=" * 90)
    print("RESUMEN")
    print("=" * 90)

    print(
        "EVALUADAS :",
        len(results),
    )

    print(
        "PROBLEMAS :",
        len(problems),
    )

    print(
        "TOTAL SEC :",
        round(total, 2),
    )

    if results:

        print(
            "AVG SEC   :",
            round(
                total / len(results),
                2,
            ),
        )

    print(
        "CSV       :",
        output_path,
    )

    print()
    print("=" * 90)

    if problems:

        print(
            "TECHNICAL RESULT: FAIL"
        )

        print(
            "No realizar auditoria semantica "
            "hasta resolver los problemas."
        )

    else:

        print(
            "TECHNICAL RESULT: PASS"
        )

        print(
            "Siguiente paso: auditoria "
            "semantica del Experimento 72."
        )

    print("=" * 90)

    return results


def main():

    parser = argparse.ArgumentParser()

    parser.add_argument(
        "input",
        type=Path,
    )

    parser.add_argument(
        "--salida",
        type=Path,
        default=Path(
            "auditoria_af72_familia_completa.csv"
        ),
    )

    parser.add_argument(
        "--model",
        default="gpt-5.6-luna",
    )

    parser.add_argument(
        "--timeout",
        type=int,
        default=600,
    )

    parser.add_argument(
        "--limit",
        type=int,
        default=None,
    )

    args = parser.parse_args()

    run_audit(
        input_path=args.input,
        output_path=args.salida,
        model=args.model,
        timeout=args.timeout,
        limit=args.limit,
    )


if __name__ == "__main__":
    main()
